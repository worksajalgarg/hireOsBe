"""Fallback parse: python-docx for DOCX; PyMuPDF (+ EasyOCR) for PDFs."""

from __future__ import annotations

import asyncio
from pathlib import Path

from .docling_parser import ParsedDocument

SPARSE_TEXT_THRESHOLD = 80


class OcrFallbackError(RuntimeError):
    """Raised when fallback parsers fail to produce usable text."""


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise OcrFallbackError(
            "python-docx is not installed. pip install python-docx"
        ) from exc

    document = Document(str(path))
    parts: list[str] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [((cell.text or "").strip()) for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n\n".join(parts).strip()


def _extract_pymupdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise OcrFallbackError(
            "pymupdf is not installed. pip install pymupdf"
        ) from exc

    parts: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text") or "")
    return "\n\n".join(parts).strip()


def _extract_easyocr(path: Path) -> str:
    try:
        import fitz
    except ImportError as exc:
        raise OcrFallbackError("pymupdf is not installed (needed for OCR page render)") from exc
    import numpy as np

    try:
        import easyocr
    except ImportError as exc:
        raise OcrFallbackError("easyocr is not installed") from exc

    reader = easyocr.Reader(["en"], gpu=False, verbose=False)
    lines: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            if pix.n == 4:
                img = img[:, :, :3]
            results = reader.readtext(img)
            page_text = "\n".join(item[1] for item in results if item[1])
            if page_text:
                lines.append(page_text)
    return "\n\n".join(lines).strip()


def _fallback_sync(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()

    if suffix == ".doc":
        try:
            from .legacy_doc import LegacyDocError, extract_text_from_legacy_doc

            text = extract_text_from_legacy_doc(path)
        except LegacyDocError as exc:
            raise OcrFallbackError(str(exc)) from exc
        except Exception as exc:
            raise OcrFallbackError(f"Legacy .doc extraction failed: {exc}") from exc
        if len(text) < SPARSE_TEXT_THRESHOLD:
            raise OcrFallbackError("Legacy .doc text extraction produced sparse/empty output")
        return ParsedDocument(markdown=text, tables=[], source="legacy_doc")

    if suffix in {".docx"}:
        try:
            text = _extract_docx(path)
        except OcrFallbackError:
            raise
        except Exception as exc:
            raise OcrFallbackError(f"DOCX extraction failed: {exc}") from exc
        if len(text) < SPARSE_TEXT_THRESHOLD:
            raise OcrFallbackError("DOCX text extraction produced sparse/empty output")
        return ParsedDocument(markdown=text, tables=[], source="python-docx")

    if suffix in {".md", ".markdown", ".adoc", ".asciidoc", ".html", ".htm"}:
        try:
            text = path.read_text(encoding="utf-8", errors="replace").strip()
        except Exception as exc:
            raise OcrFallbackError(f"Text file read failed: {exc}") from exc
        if len(text) < SPARSE_TEXT_THRESHOLD:
            raise OcrFallbackError("Text extraction produced sparse/empty output")
        return ParsedDocument(markdown=text, tables=[], source="plaintext")

    # PDF + images (and anything PyMuPDF can open)
    pymupdf_error: str | None = None
    try:
        text = _extract_pymupdf(path)
    except Exception as exc:
        text = ""
        pymupdf_error = str(exc)

    if len(text) >= SPARSE_TEXT_THRESHOLD:
        return ParsedDocument(markdown=text, tables=[], source="pymupdf")

    try:
        ocr_text = _extract_easyocr(path)
    except Exception as exc:
        detail = pymupdf_error or "sparse PyMuPDF text"
        raise OcrFallbackError(
            f"OCR fallback failed ({detail}); EasyOCR error: {exc}"
        ) from exc

    if not ocr_text.strip():
        raise OcrFallbackError("EasyOCR returned empty text")

    return ParsedDocument(markdown=ocr_text, tables=[], source="pymupdf_easyocr")


async def parse_with_ocr_fallback(path: Path) -> ParsedDocument:
    return await asyncio.to_thread(_fallback_sync, path)
