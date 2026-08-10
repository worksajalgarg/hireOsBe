from pathlib import Path

import pytest

from app.model_gateway.routing_config import load_routing_config
from app.model_gateway.use_case_policy import USE_CASE_POLICIES

_ROUTING_CONFIG_PATH = Path(__file__).resolve().parent.parent / "config" / "model_routing.yaml"


@pytest.fixture(autouse=True, scope="session")
def _load_real_routing_config() -> None:
    """USE_CASE_POLICIES starts empty (see use_case_policy.py) and is
    normally populated at process startup by worker.py/main.py. Tests that
    call get_policy()/apply_provider_priority() against real use-case names
    (e.g. test_use_case_policy_priority.py) need it populated the same way
    production is, from the real config/model_routing.yaml — not a fixture
    file, so those tests keep asserting against actual production chains."""
    USE_CASE_POLICIES.update(load_routing_config(_ROUTING_CONFIG_PATH))


@pytest.fixture(scope="session")
def synthetic_docx_bytes() -> bytes:
    """A minimal, entirely synthetic .docx generated in-memory — never a
    real resume/JD, given this project's strict candidate-PII posture (see
    docs/adr/0007-no-pydantic-ai-for-structured-extraction.md's sibling
    parsing work). Not written to disk or committed — regenerated per test
    session so there's no binary fixture to accidentally mistake for real
    data."""
    import io

    import docx

    doc = docx.Document()
    doc.add_heading("Test Candidate", level=1)
    doc.add_paragraph("Placeholder Engineer with 5 years of made-up experience.")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Example Corp — Placeholder Engineer (2021-Present): Did placeholder work.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def synthetic_pdf_bytes() -> bytes:
    """A minimal, entirely synthetic .pdf generated in-memory — same
    synthetic-only rationale as synthetic_docx_bytes above."""
    import io

    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 750, "Test Candidate")
    c.drawString(72, 730, "Placeholder Engineer with 5 years of made-up experience.")
    c.drawString(72, 710, "Experience: Example Corp - Placeholder Engineer (2021-Present)")
    c.save()
    return buf.getvalue()
